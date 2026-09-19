#!/usr/bin/env python
"""
Assembles a branded, cited .pptx deck from a synthesis JSON spec.

Usage: python generate_deck.py <input_json_path> <output_pptx_path>

Input JSON shape (see src/lib/types.ts DeckSynthesis):
{
  "deck_title": str,
  "topic": str,
  "slides": [{"title": str, "bullets": [str], "speaker_notes": str, "citation_indices": [int]}],
  "citations": [{"index": int, "title": str, "authors": [str], "year": int|null, "url": str|null}]
}
"""
import json
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Brand theme — adjust these to match your deck's actual brand.
BRAND_PRIMARY = RGBColor(0x1A, 0x1A, 0x2E)
BRAND_ACCENT = RGBColor(0x0F, 0x4C, 0x81)
BRAND_TEXT = RGBColor(0x2D, 0x2D, 0x2D)
FONT_NAME = "Calibri"


def style_title(shape, size=32, color=BRAND_PRIMARY, bold=True):
    shape.text_frame.paragraphs[0].font.size = Pt(size)
    shape.text_frame.paragraphs[0].font.bold = bold
    shape.text_frame.paragraphs[0].font.color.rgb = color
    shape.text_frame.paragraphs[0].font.name = FONT_NAME


def add_title_slide(prs, deck_title, topic):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = deck_title
    style_title(slide.shapes.title, size=40)
    if len(slide.placeholders) > 1:
        subtitle = slide.placeholders[1]
        subtitle.text = f"Research synthesis: {topic}"
        subtitle.text_frame.paragraphs[0].font.size = Pt(18)
        subtitle.text_frame.paragraphs[0].font.color.rgb = BRAND_ACCENT
    return slide


def add_content_slide(prs, slide_spec, citation_by_index):
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = slide_spec["title"]
    style_title(slide.shapes.title, size=28)

    body = slide.placeholders[1].text_frame
    body.clear()
    bullets = slide_spec.get("bullets", [])
    for i, bullet in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = BRAND_TEXT
        p.font.name = FONT_NAME

    citation_indices = slide_spec.get("citation_indices", [])
    if citation_indices:
        labels = ", ".join(f"[{i}]" for i in citation_indices if i in citation_by_index)
        if labels:
            p = body.add_paragraph()
            p.text = f"Sources: {labels}"
            p.font.size = Pt(12)
            p.font.italic = True
            p.font.color.rgb = BRAND_ACCENT
            p.font.name = FONT_NAME

    notes = slide_spec.get("speaker_notes", "")
    if notes:
        slide.notes_slide.notes_text_frame.text = notes

    return slide


def add_references_slide(prs, citations):
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "References"
    style_title(slide.shapes.title, size=28)

    body = slide.placeholders[1].text_frame
    body.clear()
    for i, citation in enumerate(citations):
        authors = ", ".join(citation.get("authors") or [])
        year = citation.get("year")
        year_str = f" ({year})" if year else ""
        url = citation.get("url")
        url_str = f" — {url}" if url else ""
        text = f"[{citation['index']}] {citation['title']}{year_str}. {authors}{url_str}"

        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = text
        p.font.size = Pt(12)
        p.font.color.rgb = BRAND_TEXT
        p.font.name = FONT_NAME


def build_deck(spec: dict, output_path: str) -> None:
    prs = Presentation()

    add_title_slide(prs, spec["deck_title"], spec.get("topic", ""))

    citation_by_index = {c["index"]: c for c in spec.get("citations", [])}
    for slide_spec in spec.get("slides", []):
        add_content_slide(prs, slide_spec, citation_by_index)

    if spec.get("citations"):
        add_references_slide(prs, spec["citations"])

    prs.save(output_path)


def main():
    if len(sys.argv) != 3:
        print("Usage: python generate_deck.py <input_json_path> <output_pptx_path>", file=sys.stderr)
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]
    with open(input_path, "r", encoding="utf-8") as f:
        spec = json.load(f)

    build_deck(spec, output_path)
    print(f"Deck written to {output_path}")


if __name__ == "__main__":
    main()
